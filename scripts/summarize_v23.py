"""Build research report, figures and compact presentation from actual outputs."""
import csv,json,statistics
from pathlib import Path
import numpy as np
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
ROOT=Path(__file__).resolve().parents[1];OUT=ROOT/'results/v23'


def load(name):
    with (OUT/(name+'.csv')).open(encoding='utf-8-sig') as f:return list(csv.DictReader(f))


def avg(rows,key):return statistics.mean(float(r[key]) for r in rows if r[key]!='')


def run():
    rows=load('lexicographic');beta=load('beta_sensitivity');stress=load('stress_repeated');paired=load('paired_statistics')
    assert len(rows)==300 and len(beta)==140 and len(stress)==180 and len(paired)==72
    assert all(int(r['violations'])==0 for r in rows+beta+stress)
    intake=json.loads((OUT/'intake_summary.json').read_text(encoding='utf-8'))
    methods=['greedy','static','rolling','lexicographic','rolling_forecast']
    lines=['# 事件链与模型接入：实际实验结果','',
        '本次沿用康定案例背景，不添加其他灾种。调度新实验与V2.2旧960组分开；旧结果用于配对bootstrap，新异构事件用于词典序和权重对照。','',
        '## 真实 DeepSeek 抽取','',
        f"模型：`{intake['model']}`。50条人工编写、人工标注的模拟文本，使用同一输入比较规则与实时API；这是工程诊断集，不是独立真实灾情测试集。规则与数据生成由同一开发过程编写，其高分存在模板贴合优势。",'',
        '| 方法 | 事件类型准确率 | 地点准确率 | 伤亡F1 | 道路F1 | JSON有效率 | 平均延迟/ms |',
        '|---|---:|---:|---:|---:|---:|---:|']
    for r in intake['results']:
        lines.append('| '+r['method']+' | '+' | '.join(f"{100*r[k]:.1f}%" for k in ['event_type_accuracy','location_accuracy','casualties_f1','roads_f1','json_valid_rate'])+f" | {r['mean_latency_ms']:.1f} |")
    live=intake['results'][1]
    lines+=['',f"实时结构与依据校验通过率为 {100*live['validated_live_rate']:.1f}%。失败调用在DeepSeek指标中按空预测处理，离线回退不会冒充模型成功。部分失败来自模型推测未给出的节点或把站点ID放进node；原始输出完整保留。事件类型与地点是整条报告的集合精确匹配；伤亡F1按地点/字段/计数模式/数值的微平均集合计算，道路按事件类型/道路ID计算。",'',
        '初轮 `intake_*_initial.json` 保留；随后修复中文字符紧邻编号时的词边界校验，再运行50条。未修改提示词和标签来追求分数。两轮都属于开发过程诊断，不能称为未知数据泛化评估。','',
        '## 词典序与异构事件对照','',
        '20种子 × 3事件强度 × 5方法 = 300组；固定每类资源50%，同类型最少保留1项。新增任务随机改变地点、资源组合、P1比例、发布时间与时长；中/高强度在32分钟加入资源故障，高强度在55分钟加入第二条封路。各方法使用相同的已生成输入及执行器。','',
        '| 方法 | P1完成率 | 已开始任务平均响应/分 | 未满足资源单位 | 改派 | 撤回 |', '|---|---:|---:|---:|---:|---:|']
    for method in methods:
        r=[r for r in rows if r['method']==method]
        lines.append(f"| {method} | {100*avg(r,'urgent_satisfaction'):.2f}% | "+' | '.join(f'{avg(r,k):.2f}' for k in ['mean_response','unmet','reassignments','withdrawals'])+' |')
    lex=[r for r in rows if r['method']=='lexicographic'];calls=sum(int(r['solver_calls']) for r in lex)
    lines+=['',f"词典序共{calls}轮调用，{sum(int(r['lex_proven_p1_calls']) for r in lex)}轮证明首层P1缺口最优，{sum(int(r['lex_complete_calls']) for r in lex)}轮完成全部四层证明。每轮总求解预算0.3秒，各层均分剩余预算；首个未证明最优的层级即停止，保留可行解。",'',
        '贪心是约束感知、优先级驱动的贪心；Static仅初始规划，其低条件响应与大量未满足并存。不要将上述均值解释为所有资源/任务条件下的优势。','',
        '## 配对统计','',
        'V2.2原960组数据按资源比例/事件强度分格，同种子取滚动减约束感知贪心。每格20对，10,000次固定种子bootstrap，报告平均差、中位差、点态95%百分位区间、配对标准化效应、胜率和平局率。胜率分母包括平局；并非只在非平局中计算。未做多重比较校正，按探索性结果解读，不用是否跨0筛选展示。', '',
        '![配对胜率](../results/v23/paired_win_rates.png)', '',
        '完整72项分格统计见 [paired_statistics.csv](../results/v23/paired_statistics.csv)。响应降低为负差、P1提高为正差；条件响应不能单独替代完成率和未满足量。','',
        '## 响应权重敏感性','', '| β | P1完成率 | 条件响应/分 | 未满足单位 |', '|---|---:|---:|---:|']
    summaries=[]
    for value in (0,1,2,5,10,20,50):
        selected=[r for r in beta if int(r['beta'])==value]
        summary=dict(beta=value,p1=avg(selected,'urgent_satisfaction'),response=avg(selected,'mean_response'),unmet=avg(selected,'unmet'));summaries.append(summary)
        lines.append(f"| {value} | {100*summary['p1']:.2f}% | {summary['response']:.2f} | {summary['unmet']:.2f} |")
    fig,axes=plt.subplots(1,2,figsize=(10,4),constrained_layout=True)
    for ax,y,label in [(axes[0],'p1','P1 completion'),(axes[1],'unmet','Unmet resource units')]:
        ax.plot([s['response'] for s in summaries],[s[y] for s in summaries],'-o')
        for s in summaries:ax.annotate(str(s['beta']),(s['response'],s[y]),xytext=(4,5),textcoords='offset points')
        ax.set(xlabel='Conditional mean response (minutes)',ylabel=label);ax.grid(alpha=.2)
    fig.suptitle('Beta sensitivity: 20 paired seeds, 50% resources, high event intensity')
    fig.savefig(OUT/'beta_tradeoffs.png',dpi=180);plt.close(fig)
    lines+=['','![权重折中](../results/v23/beta_tradeoffs.png)','',
        '每个β均用20个相同种子，不按本表自动选择“最优β”。先报告多指标折中；默认β=10仅保留兼容操作点，不声称经独立测试选优。','',
        '## 重复压力测试','', '| 情景 | 次数 | P50/ms | P95/ms | P99/ms | 违规 |', '|---|---:|---:|---:|---:|---:|']
    for name in sorted({r['scenario'] for r in stress}):
        selected=[r for r in stress if r['scenario']==name];latency=[float(r['latency_ms']) for r in selected]
        lines.append(f'| {name} | {len(selected)} | '+' | '.join(f'{np.quantile(latency,q):.2f}' for q in (.5,.95,.99))+' | 0 |')
    lines+=['','每情景重复30次，同场景输入、完整执行，每分钟约束审计。该延迟覆盖内存服务的事件处理和重规划，不含HTTP、SQLite与模型网络调用；并发运行与机器负载会影响数值，P99是小样本插值估计。压力实验沿用V2.2连续处理多事件口径；页面新批次接口另有单次重规划的原子确认测试。','',
        '300组调度、140组β实验及180组压力测试均0约束违规。九阶段页面另外验证了修正历史、次生风险、数据延迟、增援、拒绝候选、真实模型调用以及词典序切换。','',
        '## 能与不能得出的结论','',
        '本轮证明了事件和模型输入能进入经过校验、人工确认、统一重规划的执行链；统计分析更清楚地暴露不同调度方法的折中。它不证明大模型必然优于规则、词典序必然优于加权优化或真实救援有效。',
        '', '地形易发性、资源保留/完整两阶段随机优化仍未实现。本轮没有扩大灾种，也没有把模拟阶段写成历史事实。重现说明见 [README-V23.md](../README-V23.md)。']
    (ROOT/'docs/results-v23.md').write_text('\n'.join(lines)+'\n',encoding='utf-8')
    presentation(intake,rows,summaries)
    report=ROOT/'docs/results-v23.md'
    report.write_text(report.read_text(encoding='utf-8')+'\n差值热力图采用各指标统一的对称色标：\n\n![配对平均差](../results/v23/paired_differences.png)\n\n压力实验混合了各工作进程首次初始化预测器与后续缓存命中；delayed_sensor 的尾部耗时包含这种冷启动成本，不能解释为稳定热启动延迟。\n',encoding='utf-8')


def presentation(intake,rows,betas):
    from pptx import Presentation
    from pptx.util import Inches,Pt
    deck=Presentation();deck.slide_width=Inches(13.33);deck.slide_height=Inches(7.5)
    slides=[('动态灾害下的连续决策','康定案例背景 · 预测、事件理解与异构资源调度\n一个真实案例，多阶段模拟事件链\n本机研究原型；不还原实际指挥',None),
        ('问题与研究边界','动态信息会新增、延迟和修正，资源也会失效。\n目标：在约束内连续形成可核验方案。\n真实最终环境产品用于事后回放；任务与路况事件为模拟。',None),
        ('数据到执行的计算链','Bronze / Silver / Gold → 联合预测 → 区域排序\n报告 → DeepSeek / 规则 → schema与证据校验 → 人工确认\nCP-SAT + 路网 → 执行状态 → 事件重规划 → 决策追踪',None),
        ('真实历史环境与联合预测','GPM Final + ERA5-Land：三点、两变量、小时联合输入\n3,672条原生记录，2,442条小时记录；不补造水位和位移\nPersistence / Ridge / GB；固定训练、递归多时距验证\n复杂模型未全面优于Persistence。',None),
        ('一个案例：九阶段事件链','强降雨风险 → 首次灾情 → 新增失联 → 封路与延迟\n执行资源故障 → 增援 → 邻区次生风险\n信息修正与二次封路 → 趋稳恢复\n所有阶段均为实验构造，逐次人工确认。',None),
        ('DeepSeek：真实调用与证据限制',f"50条模拟文本；{intake['model']}\n事件类型准确率 {intake['results'][1]['event_type_accuracy']:.0%}；伤亡与道路F1均为1.00\n推测地点被拒绝，回退不计作模型成功。\n规则与标签模板相关，不能据此宣称真实场景泛化。",None),
        ('调度目标与词典序基线','加权：U + βR + 10L + T + λC\n词典序：P1缺口 → P1响应 → 其他缺口 → 行程与扰动\n仅固定已经证明最优的层级；超时停止并保留可行解\n共用复合资源同步、执行锁定与道路可达约束。',None),
        ('配对分析：不只报告均值','旧960组分格配对；20种子/格，10,000次bootstrap\n平均差、中位差、95%点区间、胜率、平局\n探索性分析，不做全面优势或多重检验结论。','paired_win_rates.png'),
        ('β不是随意挑一个最优值','0 / 1 / 2 / 5 / 10 / 20 / 50；相同20种子\n查看P1完成、条件响应与缺口的折中\n保留β=10作为兼容操作点，未声称独立验证选优。','beta_tradeoffs.png'),
        ('压力测试与可解释执行','6类情景 × 30次；报告P50 / P95 / P99\n300组调度 + 140组β + 180组压力：0约束违规\n内存事件处理延迟不包含HTTP、数据库或模型网络时间。',None),
        ('演示：为什么方案变化','展示模型原文与候选 → 人工确认多事件\n地图/Gantt → 改派与撤回 → ETA前后变化\n查看8→13→15人的报告版本链，核对无重复累计\n再切换词典序，查看每层证明状态。',None),
        ('研究结论与下一步','没有证明复杂模型一定胜过规则或贪心。\n价值是可复现地说明：在哪些条件、哪些指标上出现折中。\n后续：静态易发性层、资源保留与容量补救模型。\n现阶段保留单案例和清晰的模拟边界。',None)]
    notes=[]
    for index,(title,body,figure) in enumerate(slides,1):
        slide=deck.slides.add_slide(deck.slide_layouts[5]);slide.shapes.title.text=title
        box=slide.shapes.add_textbox(Inches(.6),Inches(1.3),Inches(12),Inches(1.4 if figure else 5.5))
        box.text_frame.word_wrap=True
        for i,line in enumerate(body.split('\n')):
            p=box.text_frame.paragraphs[0] if i==0 else box.text_frame.add_paragraph();p.text=line;p.font.size=Pt(20 if figure else 26)
        if figure:
            width=7.5 if figure=='paired_win_rates.png' else 10
            slide.shapes.add_picture(str(OUT/figure),Inches((13.33-width)/2),Inches(3.0),width=Inches(width))
        notes.append(f'## 第{index}页：{title}\n\n'+body.replace('\n','。')+'\n')
    folder=ROOT/'deliverables/v23';folder.mkdir(exist_ok=True)
    deck.save(folder/'research_update.pptx')
    (folder/'speaker_notes.md').write_text('# 汇报讲稿要点\n\n'+'\n'.join(notes),encoding='utf-8')


if __name__=='__main__':run()
